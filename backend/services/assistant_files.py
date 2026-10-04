"""
services/assistant_files.py

General file handling for the AI assistant.

Supported user files include:
  - PDF
  - DOCX
  - XLSX / XLSM
  - CSV / TSV
  - TXT / Markdown
  - JSON / XML
  - YAML
  - source-code / log / SQL files
  - images

The assistant does NOT require an uploaded file to be a DFS salary file.

A file is first handled as a general file. If its contents look like
DFS salary/contest data, the assistant can optionally pass it to the
specialized DFS parser/lineup engine.

Security model
  - Database access uses the USER-SCOPED client, so row-level security applies.
  - Storage access uses the service client, but every path is checked
    to start with "<user_id>/" before it is touched.
"""

from __future__ import annotations

import base64
import csv
import io
import json
import re
import uuid
from typing import Any

from config.settings import settings
from config.supabase_client import supabase as _service_client


BUCKET = "assistant-files"
SIGNED_URL_TTL = 3600
MAX_IMAGE_BYTES = 8 * 1024 * 1024


# ============================================================
# MIME TYPES
# ============================================================

IMAGE_TYPES = {
    "image/png",
    "image/jpeg",
    "image/webp",
    "image/gif",
}

TEXT_EXTENSIONS = {
    "txt",
    "md",
    "csv",
    "tsv",
    "json",
    "xml",
    "html",
    "htm",
    "yaml",
    "yml",
    "log",
    "sql",
    "py",
    "js",
    "jsx",
    "ts",
    "tsx",
    "css",
}

# Formats the assistant is allowed to generate.
GENERATED_FORMATS = {
    "csv": "text/csv",
    "txt": "text/plain",
    "md": "text/markdown",
    "json": "application/json",
    "html": "text/html",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


_LIST_COLUMNS = (
    "id,name,mime_type,size_bytes,is_generated,"
    "conversation_id,created_at"
)


# ============================================================
# ERRORS
# ============================================================

class FileError(Exception):
    """User-facing file problem."""

    def __init__(
        self,
        message: str,
        status_code: int = 400,
    ) -> None:
        super().__init__(message)
        self.message = message
        self.status_code = status_code


# ============================================================
# SMALL HELPERS
# ============================================================

def _storage():
    if _service_client is None:
        raise FileError(
            "File storage is not configured.",
            503,
        )

    return _service_client.storage.from_(BUCKET)


def _safe_name(name: str) -> str:
    cleaned = re.sub(
        r"[^\w.\-]+",
        "_",
        (name or "").strip(),
    )

    return cleaned[:120] or "file"


def _ext(name: str) -> str:
    name = (name or "").lower()

    return (
        name.rsplit(".", 1)[-1]
        if "." in name
        else ""
    )


def _valid_uuid(value: str) -> str:
    try:
        return str(uuid.UUID(str(value)))
    except (
        ValueError,
        AttributeError,
        TypeError,
    ):
        raise FileError(
            "File not found.",
            404,
        )


def _check_path(
    user_id: str,
    path: str,
) -> None:
    """
    A storage path must live inside the caller's own folder.
    """

    parts = path.split("/")

    if (
        len(parts) < 2
        or parts[0] != user_id
        or any(
            p in ("", ".", "..")
            for p in parts
        )
    ):
        raise FileError(
            "Invalid file path.",
            403,
        )


def _max_bytes() -> int:
    return (
        settings.assistant_max_upload_mb
        * 1024
        * 1024
    )


# ============================================================
# FILE TYPE INFORMATION
# ============================================================

def get_file_type(
    name: str,
    mime_type: str | None,
) -> dict[str, Any]:
    """
    Return general information about a file.

    This does NOT decide whether a file is a DFS file.
    """

    ext = _ext(name)
    mime = (
        mime_type
        or "application/octet-stream"
    ).lower()

    if ext in {"csv", "tsv"}:
        category = "tabular"
    elif ext in {"xlsx", "xlsm", "xls"}:
        category = "spreadsheet"
    elif ext in {"pdf"} or mime == "application/pdf":
        category = "pdf"
    elif ext == "docx":
        category = "document"
    elif ext in TEXT_EXTENSIONS or mime.startswith("text/"):
        category = "text"
    elif mime in IMAGE_TYPES:
        category = "image"
    elif ext == "json" or mime == "application/json":
        category = "json"
    else:
        category = "binary"

    return {
        "extension": ext,
        "mime_type": mime,
        "category": category,
        "is_image": mime in IMAGE_TYPES,
        "is_tabular": category in {
            "tabular",
            "spreadsheet",
        },
    }


# ============================================================
# DFS FILE DETECTION
# ============================================================

def detect_dfs_file(
    name: str,
    text: str | None,
) -> dict[str, Any]:
    """
    Lightweight detection for DFS salary/contest files.
    """

    ext = _ext(name)

    if ext not in {
        "csv",
        "tsv",
        "xlsx",
        "xlsm",
        "xls",
    }:
        return {
            "is_possible_dfs": False,
            "confidence": 0.0,
            "provider": None,
            "reason": "File is not a typical tabular DFS format.",
        }

    if not text:
        return {
            "is_possible_dfs": False,
            "confidence": 0.0,
            "provider": None,
            "reason": "No readable text was extracted.",
        }

    sample = text[:100_000].lower()

    dk_terms = [
        "draftkings",
        "roster position",
        "salary",
        "game info",
        "avgpointspergame",
        "avg points",
    ]

    fd_terms = [
        "fanduel",
        "nickname",
        "salary",
        "position",
        "team",
    ]

    dk_score = sum(
        1
        for term in dk_terms
        if term in sample
    )

    fd_score = sum(
        1
        for term in fd_terms
        if term in sample
    )

    salary_signal = (
        "salary" in sample
        or "salary ($)" in sample
        or "salary($)" in sample
    )

    roster_signal = (
        "roster position" in sample
        or "roster_position" in sample
    )

    if dk_score >= 2 or (
        salary_signal and roster_signal
    ):
        confidence = min(
            1.0,
            0.50 + (dk_score * 0.10),
        )

        return {
            "is_possible_dfs": True,
            "confidence": confidence,
            "provider": "draftkings",
            "reason": (
                "File contains multiple DraftKings/DFS "
                "salary or roster indicators."
            ),
        }

    if fd_score >= 2 and salary_signal:
        confidence = min(
            1.0,
            0.45 + (fd_score * 0.10),
        )

        return {
            "is_possible_dfs": True,
            "confidence": confidence,
            "provider": "fanduel",
            "reason": (
                "File contains multiple FanDuel/DFS "
                "salary indicators."
            ),
        }

    return {
        "is_possible_dfs": False,
        "confidence": 0.0,
        "provider": None,
        "reason": (
            "File does not contain enough DFS salary "
            "indicators."
        ),
    }


# ============================================================
# TEXT EXTRACTION
# ============================================================

def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    limit = settings.assistant_file_text_chars

    reader = PdfReader(
        io.BytesIO(data)
    )

    pages: list[str] = []
    size = 0

    for number, page in enumerate(
        reader.pages,
        1,
    ):
        text = page.extract_text() or ""

        pages.append(
            f"[Page {number}]\n{text}"
        )

        size += len(text)

        if size >= limit:
            break

    return "\n\n".join(pages)


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(
        io.BytesIO(data)
    )

    parts = [
        p.text
        for p in doc.paragraphs
        if p.text.strip()
    ]

    for table in doc.tables:
        for row in table.rows:
            parts.append(
                " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                )
            )

    return "\n".join(parts)


def _xlsx_text(data: bytes) -> str:
    from openpyxl import load_workbook

    limit = settings.assistant_file_text_chars

    wb = load_workbook(
        io.BytesIO(data),
        read_only=True,
        data_only=True,
    )

    lines: list[str] = []
    size = 0

    try:
        for ws in wb.worksheets:
            lines.append(
                f"[Sheet: {ws.title}]"
            )

            for i, row in enumerate(
                ws.iter_rows(values_only=True)
            ):
                if i >= 10_000:
                    break

                if size >= limit:
                    break

                line = " | ".join(
                    ""
                    if c is None
                    else str(c)
                    for c in row
                )

                lines.append(line)

                size += len(line)

            if size >= limit:
                break

    finally:
        wb.close()

    return "\n".join(lines)


def _csv_preview(
    data: bytes,
    delimiter: str = ",",
) -> str:
    limit = settings.assistant_file_text_chars

    decoded = data.decode(
        "utf-8",
        errors="replace",
    )

    buffer = io.StringIO(
        decoded,
        newline="",
    )

    reader = csv.reader(
        buffer,
        delimiter=delimiter,
    )

    lines: list[str] = []
    size = 0

    for row_number, row in enumerate(
        reader
    ):
        if row_number >= 50_000:
            break

        line = " | ".join(
            str(value)
            for value in row
        )

        if size + len(line) > limit:
            break

        lines.append(line)
        size += len(line) + 1

    return "\n".join(lines)


def _json_text(data: bytes) -> str:
    limit = settings.assistant_file_text_chars

    raw = data.decode(
        "utf-8",
        errors="replace",
    )

    try:
        parsed = json.loads(raw)

        formatted = json.dumps(
            parsed,
            indent=2,
            ensure_ascii=False,
            default=str,
        )

        return formatted[:limit]

    except Exception:
        return raw[:limit]


def extract_text(
    name: str,
    mime_type: str | None,
    data: bytes,
) -> str | None:
    """
    Return readable file content as text.

    Returns None for formats that are not supported as text.
    """

    ext = _ext(name)

    mime = (
        mime_type or ""
    ).lower()

    try:
        if (
            ext == "pdf"
            or mime == "application/pdf"
        ):
            return _pdf_text(data)

        if ext == "docx":
            return _docx_text(data)

        if ext in {
            "xlsx",
            "xlsm",
        }:
            return _xlsx_text(data)

        if ext == "csv":
            return _csv_preview(
                data,
                delimiter=",",
            )

        if ext == "tsv":
            return _csv_preview(
                data,
                delimiter="\t",
            )

        if ext == "json":
            return _json_text(data)

        if (
            ext in TEXT_EXTENSIONS
            or mime.startswith("text/")
            or mime in {
                "application/json",
                "application/xml",
            }
        ):
            return data.decode(
                "utf-8",
                errors="replace",
            )[
                : settings.assistant_file_text_chars
            ]

    except Exception:
        return None

    return None


# ============================================================
# UPLOADED FILES
# ============================================================

def register_uploaded_file(
    user_client: Any,
    user_id: str,
    *,
    storage_path: str,
    name: str,
    mime_type: str | None,
    size_bytes: int,
    conversation_id: str | None = None,
) -> dict:
    """
    Register a file that the browser already uploaded to
    Supabase Storage.

    Designed to almost never fail:
    - Large files are still registered
    - Extraction failures do not block registration
    - Download failures still allow a record to be created
    """

    _check_path(
        user_id,
        storage_path,
    )

    # Soft size check – we still register, just limit heavy work later
    if size_bytes > _max_bytes():
        # Do not raise – continue and register with limited extraction
        pass

    if conversation_id:
        conversation_id = _valid_uuid(
            conversation_id
        )

        owned = (
            user_client
            .table("ai_conversations")
            .select("id")
            .eq(
                "id",
                conversation_id,
            )
            .limit(1)
            .execute()
        )

        if not owned.data:
            raise FileError(
                "Conversation not found.",
                404,
            )

    # ---------------------------------------------------------------
    # Download from storage – never hard-fail
    # ---------------------------------------------------------------
    data = b""
    try:
        data = _storage().download(storage_path)
    except Exception:
        # Still register the file even if we cannot download it right now
        data = b""

    # Soft size limit on actual bytes
    if data and len(data) > _max_bytes():
        # Keep only the first max bytes for extraction
        data = data[: _max_bytes()]

    mime = (
        mime_type
        or "application/octet-stream"
    ).lower()

    file_type = get_file_type(
        name,
        mime,
    )

    # ---------------------------------------------------------------
    # Extract text – never fail the whole registration
    # ---------------------------------------------------------------
    text = None

    if data and mime not in IMAGE_TYPES:
        try:
            text = extract_text(
                name,
                mime,
                data,
            )
            if text:
                text = text[
                    : settings.assistant_file_text_chars
                ]
        except Exception:
            text = None

    dfs_info = detect_dfs_file(
        name,
        text,
    )

    record = {
        "user_id": user_id,
        "conversation_id": conversation_id,
        "name": name[:255],
        "mime_type": mime,
        "size_bytes": size_bytes or len(data) or 0,
        "storage_path": storage_path,
        "extracted_text": text,
        "is_generated": False,
    }

    metadata = {
        "file_type": file_type,
        "dfs_detection": dfs_info,
    }

    try:
        result = (
            user_client
            .table("ai_files")
            .insert(
                {
                    **record,
                    "metadata": metadata,
                }
            )
            .execute()
        )

    except Exception as exc:
        error_text = str(exc).lower()

        if "metadata" not in error_text:
            raise FileError(
                "Could not save the file record.",
                500,
            )

        result = (
            user_client
            .table("ai_files")
            .insert(record)
            .execute()
        )

    if not result.data:
        raise FileError(
            "Could not save the file record.",
            500,
        )

    return result.data[0]


# ============================================================
# READING / LISTING
# ============================================================

def list_files(
    user_client: Any,
    conversation_id: str | None = None,
    limit: int = 100,
) -> list[dict]:

    query = (
        user_client
        .table("ai_files")
        .select(_LIST_COLUMNS)
    )

    if conversation_id:
        query = query.eq(
            "conversation_id",
            _valid_uuid(conversation_id),
        )

    result = (
        query
        .order(
            "created_at",
            desc=True,
        )
        .limit(limit)
        .execute()
    )

    return result.data or []


def get_file(
    user_client: Any,
    file_id: str,
) -> dict:
    file_id = _valid_uuid(
        file_id
    )

    result = (
        user_client
        .table("ai_files")
        .select("*")
        .eq("id", file_id)
        .limit(1)
        .execute()
    )

    if not result.data:
        raise FileError(
            "File not found.",
            404,
        )

    return result.data[0]


def load_attachments(
    user_client: Any,
    file_ids: list[str],
) -> list[dict]:
    ids: list[str] = []

    for fid in file_ids[:10]:
        try:
            ids.append(
                str(
                    uuid.UUID(
                        str(fid)
                    )
                )
            )
        except (
            ValueError,
            TypeError,
        ):
            continue

    if not ids:
        return []

    result = (
        user_client
        .table("ai_files")
        .select("*")
        .in_("id", ids)
        .execute()
    )

    return result.data or []


# ============================================================
# IMAGES
# ============================================================

def is_image(
    row: dict,
) -> bool:
    return (
        row.get("mime_type") or ""
    ).lower() in IMAGE_TYPES


def image_data_url(
    user_id: str,
    row: dict,
) -> str | None:
    if not is_image(row):
        return None

    _check_path(
        user_id,
        row["storage_path"],
    )

    try:
        data = _storage().download(
            row["storage_path"]
        )
    except Exception:
        return None

    if len(data) > MAX_IMAGE_BYTES:
        return None

    encoded = base64.b64encode(
        data
    ).decode("ascii")

    return (
        f"data:{row['mime_type']};base64,{encoded}"
    )


# ============================================================
# DOWNLOAD / DELETE
# ============================================================

def signed_download_url(
    user_client: Any,
    user_id: str,
    file_id: str,
) -> dict:

    row = get_file(
        user_client,
        file_id,
    )

    _check_path(
        user_id,
        row["storage_path"],
    )

    try:
        res = _storage().create_signed_url(
            row["storage_path"],
            SIGNED_URL_TTL,
            options={
                "download": row["name"]
            },
        )

    except TypeError:
        res = _storage().create_signed_url(
            row["storage_path"],
            SIGNED_URL_TTL,
        )

    url = (
        res.get("signedURL")
        or res.get("signedUrl")
    )

    if not url:
        raise FileError(
            "Could not create a download link.",
            500,
        )

    if url.startswith("/"):
        url = (
            f"{settings.supabase_url}"
            f"/storage/v1{url}"
        )

    return {
        "url": url,
        "expires_in": SIGNED_URL_TTL,
    }


def delete_file(
    user_client: Any,
    user_id: str,
    file_id: str,
) -> None:

    row = get_file(
        user_client,
        file_id,
    )

    _check_path(
        user_id,
        row["storage_path"],
    )

    try:
        _storage().remove(
            [row["storage_path"]]
        )
    except Exception:
        pass

    (
        user_client
        .table("ai_files")
        .delete()
        .eq("id", row["id"])
        .execute()
    )


# ============================================================
# GENERATED FILES
# ============================================================

def _cell(value: Any) -> Any:
    if value is None:
        return ""

    if isinstance(
        value,
        (
            int,
            float,
            bool,
        ),
    ):
        return value

    text = str(value)

    if text[:1] in (
        "=",
        "@",
    ):
        return "'" + text

    return text


def _build_csv(
    rows: list[list[Any]],
) -> bytes:

    buffer = io.StringIO()

    writer = csv.writer(
        buffer
    )

    for row in rows:
        writer.writerow(
            [
                _cell(v)
                for v in row
            ]
        )

    return buffer.getvalue().encode(
        "utf-8-sig"
    )


def _build_xlsx(
    rows: list[list[Any]],
) -> bytes:

    from openpyxl import Workbook

    wb = Workbook()

    ws = wb.active

    for row in rows:
        ws.append(
            [
                _cell(v)
                for v in row
            ]
        )

    buffer = io.BytesIO()

    wb.save(buffer)

    return buffer.getvalue()


def create_generated_file(
    user_client: Any,
    user_id: str,
    *,
    filename: str,
    file_format: str,
    text: str | None = None,
    rows: list[list[Any]] | None = None,
    conversation_id: str | None = None,
) -> dict:
    fmt = (
        file_format
        or ""
    ).lower().lstrip(".")

    if fmt not in GENERATED_FORMATS:
        raise FileError(
            (
                f"Unsupported format '{file_format}'. "
                "Use one of: "
                f"{', '.join(GENERATED_FORMATS)}."
            )
        )

    if fmt == "xlsx":

        if not rows:
            raise FileError(
                "xlsx files need 'rows'."
            )

        data = _build_xlsx(
            rows
        )

    elif fmt == "csv" and rows:

        data = _build_csv(
            rows
        )

    else:

        if text is None:
            raise FileError(
                "This file type needs 'text' content."
            )

        data = text.encode(
            "utf-8"
        )

    if len(data) > _max_bytes():
        raise FileError(
            "Generated file is too large.",
            413,
        )

    name = _safe_name(
        filename
    )

    if not name.lower().endswith(
        f".{fmt}"
    ):
        name = f"{name}.{fmt}"

    path = (
        f"{user_id}/generated/"
        f"{uuid.uuid4().hex}_{name}"
    )

    mime = GENERATED_FORMATS[fmt]

    _storage().upload(
        path,
        data,
        file_options={
            "content-type": mime,
            "upsert": "true",
        },
    )

    preview = None

    if fmt != "xlsx":
        preview = (
            data
            .decode(
                "utf-8-sig",
                errors="replace",
            )
            [
                : settings.assistant_file_text_chars
            ]
        )

    if conversation_id:
        conversation_id = _valid_uuid(
            conversation_id
        )

    record = {
        "user_id": user_id,
        "conversation_id": conversation_id,
        "name": name,
        "mime_type": mime,
        "size_bytes": len(data),
        "storage_path": path,
        "extracted_text": preview,
        "is_generated": True,
    }

    try:

        result = (
            user_client
            .table("ai_files")
            .insert(record)
            .execute()
        )

        if not result.data:
            raise FileError(
                "Could not save the generated file.",
                500,
            )

    except Exception:

        try:
            _storage().remove(
                [path]
            )
        except Exception:
            pass

        raise

    return result.data[0]